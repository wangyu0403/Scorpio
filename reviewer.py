# -*- coding: utf-8 -*-
"""需求文档评审器：按完整性 / 清晰性 / 可测试性三个维度调用 LLM 评审 Word 需求，
汇总问题并导出为 Excel 报告。"""

import os
import datetime

import pandas as pd
import json_repair
from docx import Document

# 与 generator.py 保持同一套 LLM 工厂；如项目实际路径不同，改这一行即可
from llm_factory import LLMFactory


class RequirementReviewer:
    """需求文档多维评审器。"""

    # 评审维度 -> 提示词模板文件名（模板内使用 {req_text} 占位）
    DIMENSIONS = {
        "完整性": "review_completeness.md",
        "清晰性": "review_clarity.md",
        "可测试性": "review_testability.md",
    }

    # 报告固定列顺序
    REPORT_COLUMNS = ["评审维度", "问题类型", "严重程度", "问题描述", "修改建议"]

    def __init__(self, api_key):
        # 与 TestCaseGenerator 同款构造方式
        self.llm = LLMFactory(api_key=api_key)
        # prompts 目录默认取本文件同级目录下的 prompts/
        self.prompts_dir = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "prompts"
        )

    # ------------------------------------------------------------------ #
    # 文件读取
    # ------------------------------------------------------------------ #
    def read_word(self, file_path):
        """读取 .docx：拼接正文段落 + 表格单元格文本。"""
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n".join(parts)

    def read_text(self, file_path):
        """读取 .md / .txt 等纯文本文件。"""
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    # ------------------------------------------------------------------ #
    # JSON 解析（健壮容错，参考 generator.py）
    # ------------------------------------------------------------------ #
    def _parse_json_array(self, content):
        """将模型返回内容解析为问题字典列表，任何异常都安全返回空列表。"""
        if not content or not isinstance(content, str):
            return []

        text = content.strip()

        # 1) 去掉 Markdown 代码围栏
        if text.startswith("```"):
            lines = text.splitlines()
            if len(lines) >= 2:
                text = "\n".join(lines[1:])
            text = text.strip()
            if text.endswith("```"):
                text = text[:-3]
            text = text.strip()

        if not text:
            return []

        # 2) json_repair 容错解析：自动修复截断、缺括号、多余逗号等
        try:
            parsed = json_repair.loads(text)
        except Exception:
            return []

        # 3) 处理非列表结果：优先提取常见列表字段，否则把单个字典包成列表
        if isinstance(parsed, dict):
            for key in ("issues", "data", "items", "problems", "问题", "结果"):
                value = parsed.get(key)
                if isinstance(value, list):
                    parsed = value
                    break
            else:
                parsed = [parsed]

        if not isinstance(parsed, list):
            return []

        # 4) 过滤非字典项与空字典
        result = []
        for item in parsed:
            if isinstance(item, dict) and item:
                result.append(item)
        return result

    # ------------------------------------------------------------------ #
    # 主流程
    # ------------------------------------------------------------------ #
    def review(self, req_path, output_dir, req_name=None):
        # python-docx 仅支持 .docx，老式 .doc 需先另存为 .docx  基于目前时间问题，doc格式需要手动转换为docx
        ext = os.path.splitext(req_path)[1].lower()
        if ext == ".docx":
            req_text = self.read_word(req_path)
        elif ext == ".doc":
            raise ValueError(
                f"暂不支持老式 .doc 格式：{os.path.basename(req_path)}。"
                "请用 Word 另存为 .docx 后再试。"
            )
        elif ext in (".md", ".txt"):
            req_text = self.read_text(req_path)
        else:
            raise ValueError(f"不支持的需求文件格式：{ext}")

        # 需求名：默认取文件名（不含扩展名）
        if not req_name:
            req_name = os.path.splitext(os.path.basename(req_path))[0]

        all_issues = []
        for dimension, template_file in self.DIMENSIONS.items():
            template_path = os.path.join(self.prompts_dir, template_file)
            template = self.read_text(template_path)

            # 组装 Prompt
            prompt = template.replace("{req_text}", req_text)

            print(f"[评审] 维度：{dimension} ...")
            response = self.llm.chat([{"role": "user", "content": prompt}])
            issues = self._parse_json_array(response)

            # 为每条问题补充评审维度
            for issue in issues:
                issue["评审维度"] = dimension
            all_issues.extend(issues)

        # 汇总为 DataFrame，固定列顺序
        df = pd.DataFrame(all_issues)
        for col in self.REPORT_COLUMNS:
            if col not in df.columns:
                df[col] = ""
        df = df[self.REPORT_COLUMNS].fillna("")

        # 保存 Excel
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{req_name}-需求评审报告-{timestamp}.xlsx"
        output_path = os.path.join(output_dir, filename)
        df.to_excel(output_path, index=False)

        print(f"[完成] 评审报告已保存：{output_path}（共 {len(all_issues)} 条问题）")
        return output_path